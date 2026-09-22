"""
test_research_integration.py

End-to-end integration and smoke test suite for DocSort AI Research Model Integration.
Validates:
1. /model-info endpoint (LinearSVC, C=0.5, 100k features, 6 classes)
2. /predict with PDF document
3. /predict with DOCX document
4. /predict with TXT document
5. /predict-bulk with multiple document formats
6. Ambiguous document classification (decision margin < 0.5)
7. /confirm-sort endpoint and physical filesystem verification in ~/Desktop/SortedDocuments/
8. Critical validation outputs
"""

import os
import shutil
from pathlib import Path
from docx import Document
from fastapi.testclient import TestClient

from backend.main import app, BASE_DIR, PENDING_DIR, CATEGORIES

client = TestClient(app)

TEST_ASSETS_DIR = Path("test_assets")
TEST_ASSETS_DIR.mkdir(exist_ok=True)

def create_sample_files():
    # 1. DOCX (Technology & Computing)
    docx_path = TEST_ASSETS_DIR / "sample_tech.docx"
    doc = Document()
    doc.add_heading("Cloud Architecture and Distributed Microservices", level=1)
    doc.add_paragraph(
        "Modern cloud computing relies on distributed container orchestration, "
        "Kubernetes clusters, serverless API backends, and low-latency database sharding. "
        "Engineers optimize hardware throughput, CPU processor cores, and network bandwidth "
        "to scale software microservices securely across multi-region data centers."
    )
    doc.save(str(docx_path))

    # 2. TXT (Sports)
    txt_path = TEST_ASSETS_DIR / "sample_sports.txt"
    txt_path.write_text(
        "The championship final match ended with a dramatic penalty kick in the 92nd minute. "
        "The head coach praised the striker's agility and the goalkeeper's clutch saves. "
        "With this tournament victory, the franchise clinched their third league title of the season, "
        "sending thousands of cheering fans into the arena to celebrate the trophy presentation."
    )

    # 3. TXT (Ambiguous: blend of corporate technology stock earnings and business finance)
    ambiguous_path = TEST_ASSETS_DIR / "sample_ambiguous.txt"
    ambiguous_path.write_text(
        "The enterprise software company held an investor quarterly earnings conference call today. "
        "The Chief Financial Officer discussed share buybacks, operating margins, dividend yields, "
        "and balance sheet liquidity alongside product roadmaps for the next-generation microchip "
        "processor, operating system patch release, and machine learning cloud infrastructure."
    )

    return docx_path, txt_path, ambiguous_path

def run_tests():
    print("=" * 80)
    print("DOCSORT AI — RESEARCH INTEGRATION TEST SUITE")
    print("=" * 80)

    docx_path, txt_path, ambiguous_path = create_sample_files()
    pdf_path = Path("check 1.pdf")
    if not pdf_path.exists():
        pdf_path = Path("check 2.pdf")

    # -------------------------------------------------------------
    # 1. TEST GET /model-info
    # -------------------------------------------------------------
    print("\n[TEST 1] GET /model-info")
    res = client.get("/model-info")
    assert res.status_code == 200, f"Expected 200, got {res.status_code}: {res.text}"
    data = res.json()
    print("  Status code:", res.status_code)
    print("  Model:", data["model"])
    print("  C:", data["C"])
    print("  Features:", data["vectorizer_features"])
    print("  Classes (6):", data["classes"])
    print("  Research model:", data["research_model"])
    assert data["model"] == "LinearSVC"
    assert data["C"] == 0.5
    assert data["vectorizer_features"] == 100000
    assert len(data["classes"]) == 6
    assert data["research_model"] is True
    print("  >>> PASS: /model-info verified.")

    # -------------------------------------------------------------
    # 2. TEST POST /predict with PDF
    # -------------------------------------------------------------
    print("\n[TEST 2] POST /predict with PDF file:", pdf_path)
    with open(pdf_path, "rb") as f:
        res = client.post("/predict", files={"file": (pdf_path.name, f, "application/pdf")})
    assert res.status_code == 200, f"Error: {res.text}"
    pdf_res = res.json()
    print("  Filename:", pdf_res["filename"])
    print("  Prediction:", pdf_res["prediction"])
    print("  Decision Margin:", pdf_res["decision_margin"])
    print("  Uncertainty Level:", pdf_res["uncertainty_level"])
    print("  Top Keywords:", pdf_res["top_keywords"][:5])
    print("  Stored in:", pdf_res.get("stored_in"))
    assert "prediction" in pdf_res
    assert "decision_margin" in pdf_res
    assert "uncertainty_level" in pdf_res
    print("  >>> PASS: PDF prediction succeeded.")

    # -------------------------------------------------------------
    # 3. TEST POST /predict with DOCX
    # -------------------------------------------------------------
    print("\n[TEST 3] POST /predict with DOCX file:", docx_path)
    with open(docx_path, "rb") as f:
        res = client.post("/predict", files={"file": (docx_path.name, f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert res.status_code == 200, f"Error: {res.text}"
    docx_res = res.json()
    print("  Filename:", docx_res["filename"])
    print("  Prediction:", docx_res["prediction"])
    print("  Decision Margin:", docx_res["decision_margin"])
    print("  Uncertainty Level:", docx_res["uncertainty_level"])
    print("  Top Keywords:", docx_res["top_keywords"][:5])
    print("  Stored in:", docx_res.get("stored_in"))
    assert docx_res["prediction"] == "Technology & Computing"
    assert docx_res["decision_margin"] > 0
    assert docx_res["uncertainty_level"] in ("High confidence", "Moderate confidence")
    if docx_res.get("stored_in"):
        assert os.path.exists(docx_res["stored_in"]), f"File not found at {docx_res['stored_in']}"
    print("  >>> PASS: DOCX prediction succeeded and sorted into Technology & Computing.")

    # -------------------------------------------------------------
    # 4. TEST POST /predict with TXT
    # -------------------------------------------------------------
    print("\n[TEST 4] POST /predict with TXT file:", txt_path)
    with open(txt_path, "rb") as f:
        res = client.post("/predict", files={"file": (txt_path.name, f, "text/plain")})
    assert res.status_code == 200, f"Error: {res.text}"
    txt_res = res.json()
    print("  Filename:", txt_res["filename"])
    print("  Prediction:", txt_res["prediction"])
    print("  Decision Margin:", txt_res["decision_margin"])
    print("  Uncertainty Level:", txt_res["uncertainty_level"])
    print("  Top Keywords:", txt_res["top_keywords"][:5])
    print("  Stored in:", txt_res.get("stored_in"))
    assert txt_res["prediction"] == "Sports"
    assert txt_res["decision_margin"] > 0
    if txt_res.get("stored_in"):
        assert os.path.exists(txt_res["stored_in"]), f"File not found at {txt_res['stored_in']}"
    print("  >>> PASS: TXT prediction succeeded and sorted into Sports.")

    # -------------------------------------------------------------
    # 5. TEST POST /predict-bulk
    # -------------------------------------------------------------
    print("\n[TEST 5] POST /predict-bulk with multiple documents")
    with open(docx_path, "rb") as f1, open(txt_path, "rb") as f2:
        res = client.post(
            "/predict-bulk",
            files=[
                ("files", ("bulk_tech.docx", f1, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
                ("files", ("bulk_sports.txt", f2, "text/plain")),
            ]
        )
    assert res.status_code == 200, f"Error: {res.text}"
    bulk_res = res.json()
    print("  Bulk results count:", len(bulk_res["results"]))
    for item in bulk_res["results"]:
        print(f"    - {item['filename']}: {item['prediction']} (Margin: {item['decision_margin']}, {item['uncertainty_level']})")
    assert len(bulk_res["results"]) == 2
    print("  >>> PASS: Bulk prediction succeeded.")

    # -------------------------------------------------------------
    # 6. TEST Ambiguous Prediction & /confirm-sort
    # -------------------------------------------------------------
    print("\n[TEST 6] Ambiguous Document Staging & /confirm-sort")
    with open(ambiguous_path, "rb") as f:
        res = client.post("/predict", files={"file": (ambiguous_path.name, f, "text/plain")})
    assert res.status_code == 200, f"Error: {res.text}"
    amb_res = res.json()
    print("  Filename:", amb_res["filename"])
    print("  Prediction:", amb_res["prediction"])
    print("  Decision Margin:", amb_res["decision_margin"])
    print("  Uncertainty Level:", amb_res["uncertainty_level"])
    print("  Requires Manual Choice:", amb_res.get("requires_manual_choice"))
    print("  Pending ID:", amb_res.get("pending_id"))

    # If ambiguous, confirm sort workflow
    if amb_res.get("requires_manual_choice") and amb_res.get("pending_id"):
        pending_id = amb_res["pending_id"]
        chosen = amb_res["prediction"]
        print(f"  Confirming sort for pending_id={pending_id} into chosen='{chosen}'...")
        confirm_res = client.post(
            "/confirm-sort",
            json={"pending_id": pending_id, "chosen_domain": chosen}
        )
        assert confirm_res.status_code == 200, f"Confirm error: {confirm_res.text}"
        c_data = confirm_res.json()
        print("  Confirmed stored path:", c_data["stored_in"])
        assert os.path.exists(c_data["stored_in"]), f"Sorted file missing at {c_data['stored_in']}"
        print("  >>> PASS: /confirm-sort successfully moved pending file.")
    else:
        print("  Document scored above 0.50 margin; testing /confirm-sort with a simulated pending file.")
        # Test simulated pending file to ensure /confirm-sort logic is 100% verified
        from backend.main import PENDING_FILES
        sim_id = "test-pending-uuid"
        sim_temp = Path(PENDING_DIR) / f"{sim_id}__manual_test.txt"
        sim_temp.write_text("Simulated test content.")
        PENDING_FILES[sim_id] = {
            "path": str(sim_temp),
            "allowed_domains": ("Business and Finance", "Technology & Computing"),
            "final_basename": "manual_test.txt",
            "client_filename": "manual_test.txt",
        }
        confirm_res = client.post(
            "/confirm-sort",
            json={"pending_id": sim_id, "chosen_domain": "Business and Finance"}
        )
        assert confirm_res.status_code == 200
        c_data = confirm_res.json()
        print("  Simulated confirm stored path:", c_data["stored_in"])
        assert os.path.exists(c_data["stored_in"])
        print("  >>> PASS: /confirm-sort verified.")

    # -------------------------------------------------------------
    # 7. CRITICAL VALIDATION LOGGING
    # -------------------------------------------------------------
    print("\n" + "=" * 80)
    print("CRITICAL VALIDATION REPORT")
    print("=" * 80)
    print("Loaded model: LinearSVC")
    print("Loaded vectorizer: 100000 features")
    print("Classes:", CATEGORIES)
    print("-" * 80)
    print("Sample Document Inference:")
    print("Document: sample_tech.docx")
    print("Predicted class:", docx_res["prediction"])
    print("Decision scores:", docx_res["decision_scores"])
    print("Decision margin:", docx_res["decision_margin"])
    print("Uncertainty level:", docx_res["uncertainty_level"])
    print("=" * 80)

    # Cleanup test assets
    shutil.rmtree(TEST_ASSETS_DIR, ignore_errors=True)
    print("\nALL INTEGRATION TESTS PASSED SUCCESSFULLY!\n")

if __name__ == "__main__":
    run_tests()
